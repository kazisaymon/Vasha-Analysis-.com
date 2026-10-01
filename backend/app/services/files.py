"""Turn uploaded files into Claude content blocks (images/PDF natively, others as extracted text)."""
import base64
import io

import pandas as pd
from fastapi import HTTPException, UploadFile

from ..config import settings
from . import speech

IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
TEXT_EXT = (".txt", ".md", ".json", ".xml", ".html", ".log", ".py", ".js", ".ts", ".sql", ".yaml", ".yml")


def _clip(text: str) -> str:
    if len(text) > settings.max_text_chars:
        return text[: settings.max_text_chars] + "\n...[truncated]"
    return text


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _table_text(data: bytes, name: str) -> str:
    if name.endswith(".csv"):
        sheets = {"csv": pd.read_csv(io.BytesIO(data))}
    else:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    out = []
    for sheet, df in sheets.items():
        out.append(f"[Sheet: {sheet}] {len(df)} rows x {len(df.columns)} columns")
        out.append("Columns: " + ", ".join(map(str, df.columns)))
        out.append(df.head(200).to_csv(index=False))
        try:
            out.append("Summary statistics:\n" + df.describe(include="all").T.head(40).to_string())
        except Exception:
            pass
    return "\n".join(out)


async def to_blocks(uploads: list[UploadFile]) -> list[dict]:
    if len(uploads) > settings.max_files:
        raise HTTPException(413, f"Maximum {settings.max_files} files per message.")
    blocks: list[dict] = []
    for up in uploads:
        data = await up.read()
        if len(data) > settings.max_file_bytes:
            raise HTTPException(413, f"{up.filename} is larger than {settings.max_file_bytes // 1024 // 1024} MB.")
        name = (up.filename or "file").lower()
        ctype = (up.content_type or "").lower()

        if ctype in IMAGE_TYPES:
            blocks.append({"type": "text", "text": f"[Image: {up.filename}]"})
            blocks.append({"type": "image", "source": {"type": "base64", "media_type": ctype,
                                                       "data": base64.b64encode(data).decode()}})
        elif ctype == "application/pdf" or name.endswith(".pdf"):
            blocks.append({"type": "document", "title": up.filename,
                           "source": {"type": "base64", "media_type": "application/pdf",
                                      "data": base64.b64encode(data).decode()}})
        elif ctype.startswith("audio/") or name.endswith((".mp3", ".wav", ".m4a", ".ogg", ".webm", ".flac")):
            text = await speech.transcribe(data, up.filename or "audio.webm", None)
            blocks.append({"type": "text", "text": f"[Audio transcript of {up.filename}]\n{_clip(text)}"})
        elif name.endswith(".docx"):
            blocks.append({"type": "text", "text": f"[File: {up.filename}]\n{_clip(_docx_text(data))}"})
        elif name.endswith((".csv", ".xlsx", ".xls")):
            blocks.append({"type": "text", "text": f"[Table file: {up.filename}]\n{_clip(_table_text(data, name))}"})
        elif name.endswith(TEXT_EXT) or ctype.startswith("text/"):
            blocks.append({"type": "text", "text": f"[File: {up.filename}]\n{_clip(data.decode('utf-8', 'replace'))}"})
        else:
            raise HTTPException(415, f"Unsupported file type: {up.filename}")
    return blocks
